import json
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models.job_source import (
    JobSource,
)


ALLOWED_CONNECTORS = {
    "amazon",
    "microsoft",
    "greenhouse",
    "lever",
}


def _to_bool(value) -> bool:

    if isinstance(value, bool):
        return value

    if value is None:
        return True

    return str(value).strip().lower() in {
        "true",
        "1",
        "yes",
        "y",
    }


def import_job_sources_excel(
    db: Session,
    content: bytes,
) -> dict:

    workbook = load_workbook(
        BytesIO(content),
        read_only=True,
        data_only=True,
    )

    sheet = workbook.active

    rows = list(
        sheet.iter_rows(
            values_only=True
        )
    )

    if not rows:
        raise ValueError(
            "Excel file is empty"
        )

    headers = [
        str(value)
        .strip()
        .lower()
        if value is not None
        else ""
        for value in rows[0]
    ]

    required = {
        "company_name",
        "careers_url",
        "connector_type",
    }

    missing = (
        required - set(headers)
    )

    if missing:
        raise ValueError(
            "Missing required columns: "
            + ", ".join(
                sorted(missing)
            )
        )

    imported = 0
    updated = 0
    skipped = 0

    errors = []

    for excel_row_number, row in enumerate(
        rows[1:],
        start=2,
    ):

        row_data = dict(
            zip(
                headers,
                row,
            )
        )

        if not any(
            value is not None
            for value in row
        ):
            continue

        try:
            company_name = str(
                row_data.get(
                    "company_name"
                )
                or ""
            ).strip()

            careers_url = str(
                row_data.get(
                    "careers_url"
                )
                or ""
            ).strip()

            connector_type = str(
                row_data.get(
                    "connector_type"
                )
                or ""
            ).strip().lower()

            connector_key_value = (
                row_data.get(
                    "connector_key"
                )
            )

            connector_key = (
                str(
                    connector_key_value
                ).strip()
                if connector_key_value
                is not None
                else None
            )

            if not company_name:
                raise ValueError(
                    "company_name required"
                )

            if not careers_url:
                raise ValueError(
                    "careers_url required"
                )

            if (
                connector_type
                not in
                ALLOWED_CONNECTORS
            ):
                raise ValueError(
                    "Unsupported "
                    "connector_type: "
                    f"{connector_type}"
                )

            raw_config = (
                row_data.get(
                    "config_json"
                )
            )

            if (
                raw_config is None
                or str(
                    raw_config
                ).strip() == ""
            ):
                config_json = {}
            elif isinstance(
                raw_config,
                dict,
            ):
                config_json = raw_config
            else:
                config_json = (
                    json.loads(
                        str(
                            raw_config
                        )
                    )
                )

            active = _to_bool(
                row_data.get(
                    "active"
                )
            )

            existing = (
                db.query(
                    JobSource
                )
                .filter(
                    JobSource.company_name
                    == company_name,
                    JobSource.connector_type
                    == connector_type,
                    JobSource.connector_key
                    == connector_key,
                )
                .first()
            )

            if existing:

                existing.careers_url = (
                    careers_url
                )
                existing.config_json = (
                    config_json
                )
                existing.is_active = (
                    active
                )

                updated += 1

            else:

                source = JobSource(
                    company_name=(
                        company_name
                    ),
                    careers_url=(
                        careers_url
                    ),
                    connector_type=(
                        connector_type
                    ),
                    connector_key=(
                        connector_key
                    ),
                    config_json=(
                        config_json
                    ),
                    is_active=active,
                )

                db.add(source)

                imported += 1

        except Exception as exc:

            skipped += 1

            errors.append(
                {
                    "row": (
                        excel_row_number
                    ),
                    "error": str(exc),
                }
            )

    db.commit()

    return {
        "created": imported,
        "updated": updated,
        "skipped": skipped,
        "errors": errors,
    }